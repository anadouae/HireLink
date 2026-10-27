# -*- coding: utf-8 -*-
"""Génère le pack Séance 3 + 4 : rapport, guide d'utilisation, document de tests."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]  # docs/
OUT = Path(__file__).resolve().parent
LOGO_ENSA = ROOT / "logo-ensa.png"
LOGO_UAE = ROOT / "logo-uae.png"
LOGO_HL = ROOT / "logo-hirelink-full.png"
if not LOGO_HL.exists():
    LOGO_HL = ROOT / "logo-hirelink.png"
DIAG = ROOT / "diagrammes" / "diagramme_classes_hirelink_FINAL.png"
if not DIAG.exists():
    DIAG = ROOT / "diagrammes" / "diagramme_classes_hirelink.png"

TODAY = date.today().strftime("%d/%m/%Y")


def set_run(run, *, bold=False, size=11, color=None, font="Calibri"):
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, bold=True, size=14 if level == 1 else 12, color=(122, 32, 50))
    return p


def add_p(doc, text, *, bold=False, size=11, center=False, space_after=6):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run(r, bold=True, size=10, color=(255, 255, 255))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "7A2032")
        shading.set(qn("w:val"), "clear")
        cell._teProp = cell._tc.get_or_add_tcPr()
        cell._tc.get_or_add_tcPr().append(shading)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_run(r, size=10)
    if col_widths:
        for row in t.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def cover_page(doc, title, subtitle):
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # logos row
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_UAE.exists():
        p.add_run().add_picture(str(LOGO_UAE), width=Cm(2.8))
        p.add_run("     ")
    if LOGO_ENSA.exists():
        p.add_run().add_picture(str(LOGO_ENSA), width=Cm(2.8))
        p.add_run("     ")
    if LOGO_HL.exists():
        p.add_run().add_picture(str(LOGO_HL), width=Cm(4.2))

    add_p(doc, "Université Abdelmalek Essaâdi", center=True, bold=True, size=12, space_after=2)
    add_p(doc, "École Nationale des Sciences Appliquées — Tétouan", center=True, bold=True, size=11, space_after=2)
    add_p(doc, "Filière Génie Informatique (GI)", center=True, size=11, space_after=14)

    add_p(doc, "Module : Frameworks & Technologies Web", center=True, bold=True, size=12, space_after=4)
    add_p(doc, title, center=True, bold=True, size=18, space_after=6)
    add_p(doc, subtitle, center=True, size=12, space_after=16)

    add_table(
        doc,
        ["Champ", "Valeur"],
        [
            ["Projet", "HireLink — ATS (Applicant Tracking System)"],
            ["Étudiante", "Douae El Assal"],
            ["Encadrant", "À compléter"],
            ["Séances", "Séance 3 (JPA/Hibernate) + Séance 4 (CRUD + API REST)"],
            ["Stack", "Spring Boot 3 · Java 21 · JPA/Hibernate · MySQL/MariaDB · HTML/JS · Postman"],
            ["API", "http://localhost:8081/api/v1"],
            ["Frontend", "http://localhost:5500"],
            ["Base MySQL", "hirelink (localhost:3306)"],
            ["Date", TODAY],
        ],
        col_widths=[4, 12],
    )

    add_p(
        doc,
        "Document livrable pour suivi pédagogique. Contenu aligné sur le code réellement livré.",
        center=True,
        size=9,
        space_after=8,
    )
    doc.add_page_break()


def build_rapport():
    doc = Document()
    cover_page(
        doc,
        "RAPPORT — Séances 3 & 4",
        "Modélisation JPA/Hibernate · CRUD · API REST · Postman · Interfaces",
    )

    add_heading(doc, "1. Ce qui est demandé", 1)
    add_heading(doc, "1.1 Séance 3 — Modélisation JPA / Hibernate", 2)
    add_p(doc, "Objectif : construire le modèle de données et le relier au backend.")
    add_table(
        doc,
        ["Livrable demandé", "Statut", "Preuve"],
        [
            ["Entités JPA annotées", "FAIT", "backend/.../model/*.java"],
            ["Relations (@ManyToOne, @OneToMany, @ManyToMany)", "FAIT", "User, JobOffer, Application, Tag, Membership"],
            ["Configuration Hibernate", "FAIT", "ddl-auto=update dans application.properties"],
            ["Base MySQL fonctionnelle", "FAIT", "phpMyAdmin → schéma hirelink"],
            ["Diagramme de classes", "FAIT", "docs/diagrammes/diagramme_classes_hirelink_FINAL.png"],
            ["Formulaires frontend alignés", "FAIT", "frontend/app/*"],
        ],
    )

    add_heading(doc, "1.2 Séance 4 — CRUD complet + API REST", 2)
    add_p(doc, "Objectif : CRUD backend (≥ 3 entités), API REST, interfaces, tests Postman.")
    add_table(
        doc,
        ["Livrable demandé", "Statut", "Preuve"],
        [
            ["CRUD JobOffer", "FAIT", "POST/GET/PUT/DELETE /api/v1/offers"],
            ["CRUD RefusalReason", "FAIT", "GET/POST/PUT/DELETE /api/v1/refusal-reasons"],
            ["CRUD InAppNotification", "FAIT", "POST/GET/read/DELETE /api/v1/notifications"],
            ["Couches Controller→Service→Repository", "FAIT", "packages controller/service/repository"],
            ["Frontend CRUD offres", "FAIT", "recruiter-offers.html + offer-wizard.js"],
            ["Frontend CRUD motifs", "FAIT", "admin-refusal-reasons.html"],
            ["Collection Postman", "FAIT", "docs/postman/HireLink_CRUD_Seance4.postman_collection.json"],
            ["Exécution des tests API", "FAIT", "voir DOCUMENT_TESTS + résultats Newman"],
        ],
    )

    add_heading(doc, "2. Comment les tables sont créées (point critique oral)", 1)
    add_p(
        doc,
        "Phrase à retenir : « Les tables ne sont pas créées par un script .sql artisanal. "
        "La source de vérité est le modèle Java (@Entity). Au démarrage, Hibernate lit les entities "
        "et, avec spring.jpa.hibernate.ddl-auto=update, crée ou met à jour les tables MySQL. "
        "phpMyAdmin sert uniquement à vérifier le résultat. »",
        bold=True,
    )
    add_p(doc, "Étapes exactes :")
    for i, t in enumerate(
        [
            "On écrit une classe Java annotée @Entity (ex. JobOffer.java).",
            "On démarre Spring Boot (HireLinkApplication / jar).",
            "Le conteneur Spring charge JPA/Hibernate.",
            "Hibernate compare le modèle Java et le schéma MySQL, puis applique les créations/alter.",
            "On ouvre phpMyAdmin → base hirelink → on voit job_offers, users, applications…",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_p(
        doc,
        "À NE PAS dire : « J’ai lancé le fichier schema.sql pour créer les tables. » "
        "Le fichier database/schema.sql éventuel n’est PAS la source de vérité du projet actuel.",
        bold=True,
    )

    add_heading(doc, "3. Ce que nous avons fait exactement (Séance 3)", 1)
    add_table(
        doc,
        ["Étape", "Action réalisée", "Résultat"],
        [
            ["1", "Identifier les entités métier ATS", "User, Organization, Membership, JobOffer, Application, Tag, RefusalReason, Notification, Scorecard, Matching…"],
            ["2", "Annoter les classes JPA", "package ma.ensat.hirelink.model"],
            ["3", "Mapper les relations", "OneToMany / ManyToOne / ManyToMany (job_offer_tags) / Membership"],
            ["4", "Configurer la datasource + ddl-auto=update", "application.properties"],
            ["5", "Vérifier les tables générées", "phpMyAdmin base hirelink"],
            ["6", "Produire le diagramme de classes", "PNG + PlantUML dans docs/diagrammes/"],
            ["7", "Aligner les formulaires UI", "offres, profil, pipeline, motifs"],
        ],
    )

    if DIAG.exists():
        add_p(doc, "Diagramme de classes (extrait du livrable) :", bold=True)
        doc.add_picture(str(DIAG), width=Cm(16))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "4. Ce que nous avons fait exactement (Séance 4)", 1)
    add_heading(doc, "4.1 Architecture d’une feature CRUD", 2)
    add_p(doc, "Frontend (fetch + JWT) → Controller → Service → Repository → Entity → Table MySQL")

    add_heading(doc, "4.2 Entité 1 — JobOffer", 2)
    add_table(
        doc,
        ["Opération", "HTTP", "Endpoint"],
        [
            ["Create", "POST", "/api/v1/offers"],
            ["Read list", "GET", "/api/v1/offers/mine"],
            ["Read one", "GET", "/api/v1/offers/{id}"],
            ["Update", "PUT", "/api/v1/offers/{id}"],
            ["Delete", "DELETE", "/api/v1/offers/{id}"],
        ],
    )
    add_p(doc, "Fichiers : JobOfferController · JobOfferService · JobOfferRepository · JobOffer")
    add_p(doc, "UI : frontend/app/recruiter-offers.html (créer / modifier / supprimer)")

    add_heading(doc, "4.3 Entité 2 — RefusalReason", 2)
    add_table(
        doc,
        ["Opération", "HTTP", "Endpoint"],
        [
            ["Read actifs", "GET", "/api/v1/refusal-reasons"],
            ["Read one", "GET", "/api/v1/refusal-reasons/{id}"],
            ["Create", "POST", "/api/v1/refusal-reasons (ADMIN)"],
            ["Update", "PUT", "/api/v1/refusal-reasons/{id} (ADMIN)"],
            ["Delete", "DELETE", "/api/v1/refusal-reasons/{id} (ADMIN)"],
        ],
    )
    add_p(doc, "UI : frontend/app/admin-refusal-reasons.html")

    add_heading(doc, "4.4 Entité 3 — InAppNotification", 2)
    add_table(
        doc,
        ["Opération", "HTTP", "Endpoint"],
        [
            ["Create", "POST", "/api/v1/notifications"],
            ["Read", "GET", "/api/v1/notifications"],
            ["Update (lu)", "POST", "/api/v1/notifications/{id}/read"],
            ["Delete", "DELETE", "/api/v1/notifications/{id}"],
        ],
    )

    add_heading(doc, "4.5 Postman", 2)
    add_p(doc, "Collection : docs/postman/HireLink_CRUD_Seance4.postman_collection.json")
    add_p(
        doc,
        "Ordre d’exécution : Login recruteur → CRUD Offres → Login admin → CRUD Motifs → CRUD Notifications.",
    )

    add_heading(doc, "5. Compétences couvertes", 1)
    add_p(doc, "JPA, Hibernate, ORM, relations, Spring Data, Repository, Service, REST Controller, CRUD HTTP, JWT, Postman, UI fetch.")

    add_heading(doc, "6. Phrases orales anti-surprise", 1)
    for t in [
        "Les tables viennent des @Entity + Hibernate ddl-auto=update, pas d’un .sql manuel comme source.",
        "Une feature = Controller → Service → Repository → Entity.",
        "Le frontend n’écrit jamais directement en base : il appelle l’API avec un JWT.",
        "Séance 3 = modèle ; Séance 4 = exposer et manipuler ce modèle via CRUD REST + UI + Postman.",
        "Nous avons trois CRUD complets : Offres, Motifs de refus, Notifications.",
    ]:
        add_p(doc, "• " + t, space_after=3)

    add_heading(doc, "7. Annexes — chemins", 1)
    add_p(
        doc,
        "HireLink/backend/src/main/java/ma/ensat/hirelink/{model,repository,service,controller}\n"
        "HireLink/frontend/app/\n"
        "HireLink/docs/diagrammes/\n"
        "HireLink/docs/postman/\n"
        "HireLink/docs/livraison-seance3-4/",
    )

    path = OUT / "01_RAPPORT_SEANCE3_SEANCE4_HireLink.docx"
    doc.save(path)
    return path


def build_guide():
    doc = Document()
    cover_page(
        doc,
        "GUIDE D’UTILISATION & COMPTES DE TEST",
        "Suivre chaque étape dans l’ordre — aucune étape cachée",
    )

    add_heading(doc, "0. Avant de commencer (une seule fois)", 1)
    add_p(doc, "Objectif : avoir MySQL + API + Frontend allumés.", bold=True)
    add_p(doc, "Étape A — Démarrer MySQL (XAMPP)")
    for t in [
        "Ouvrir le panneau de contrôle XAMPP.",
        "Cliquer sur Start à côté de MySQL.",
        "Vérifier que MySQL est en vert (port 3306).",
    ]:
        add_p(doc, "→ " + t, space_after=2)

    add_p(doc, "Étape B — Démarrer l’API HireLink")
    for t in [
        "Ouvrir un terminal dans le dossier : C:\\Users\\daurinia\\Desktop\\HireLink\\backend",
        "Exécuter exactement :  java -jar target\\hirelink-1.0.0-SNAPSHOT.jar",
        "Attendre le message : Started HireLinkApplication",
        "Vérifier dans le navigateur : http://localhost:8081/api/v1/health",
        "Résultat attendu : status = UP",
    ]:
        add_p(doc, "→ " + t, space_after=2)

    add_p(doc, "Étape C — Démarrer le Frontend")
    for t in [
        "Ouvrir un NOUVEAU terminal dans : C:\\Users\\daurinia\\Desktop\\HireLink\\frontend",
        "Exécuter exactement :  python -m http.server 5500",
        "Ouvrir le navigateur : http://localhost:5500/auth.html",
    ]:
        add_p(doc, "→ " + t, space_after=2)

    add_heading(doc, "1. Comptes de test (copier-coller)", 1)
    add_table(
        doc,
        ["Rôle", "Email", "Mot de passe", "Où se connecter"],
        [
            ["Administrateur", "admin@hirelink.ma", "Admin123!", "auth.html → puis app/admin.html"],
            ["Recruteur (Owner)", "recruteur@hirelink.ma", "Recruteur123!", "auth.html → app/recruiter.html"],
            ["Recruteur (Atlas)", "n.bennani@atlasdigital.ma", "Recruteur123!", "auth.html → app/recruiter.html"],
            ["Candidat", "candidat@hirelink.ma", "Candidat123!", "auth.html → app/candidate.html"],
            ["Candidat (Sara)", "s.elamrani@gmail.com", "Candidat123!", "auth.html → app/candidate.html"],
        ],
    )
    add_p(doc, "Important : utilisez exactement ces emails et mots de passe. Ne changez rien.", bold=True)

    add_heading(doc, "2. Scénario A — Connexion recruteur (2 minutes)", 1)
    for i, t in enumerate(
        [
            "Aller sur http://localhost:5500/auth.html",
            "Choisir la connexion par mot de passe (pas OTP).",
            "Email : recruteur@hirelink.ma",
            "Mot de passe : Recruteur123!",
            "Cliquer sur Connexion.",
            "Vous devez arriver sur : http://localhost:5500/app/recruiter.html",
            "Si une page set-password apparaît : créer un mot de passe puis revenir à auth.html et se reconnecter.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "3. Scénario B — CRUD Offre (Séance 4, UI)", 1)
    for i, t in enumerate(
        [
            "Depuis l’espace recruteur, cliquer sur « Mes offres » (menu gauche).",
            "URL attendue : .../app/recruiter-offers.html",
            "Cliquer sur le bouton pour créer une offre (Nouveau / Créer).",
            "Remplir : Titre = « Test CRUD Offre », Compétences = « Java,Spring », Lieu = « Tétouan », SLA = 7.",
            "Enregistrer.",
            "Vérifier que l’offre apparaît dans la liste.",
            "Ouvrir Modifier, changer le titre en « Test CRUD Offre MODIFIÉ », Enregistrer.",
            "Vérifier le nouveau titre dans la liste.",
            "Supprimer l’offre de test.",
            "Vérifier qu’elle disparaît de la liste.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "4. Scénario C — Pipeline (vue ATS)", 1)
    for i, t in enumerate(
        [
            "Menu gauche → Pipeline.",
            "URL : .../app/recruiter-pipeline.html",
            "Vous devez voir des colonnes : Soumises, Présélection, Shortlist, Entretien, Offre.",
            "Ouvrir une carte → Dossier.",
            "Pour avancer : bouton Avancer → cocher les 3 cases → Confirmer.",
            "Pour refuser : bouton Refuser → choisir un motif du catalogue → cocher les 3 cases → Confirmer.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "5. Scénario D — Admin + motifs de refus", 1)
    for i, t in enumerate(
        [
            "Se déconnecter.",
            "Se connecter avec admin@hirelink.ma / Admin123!",
            "Aller sur .../app/admin.html (supervision).",
            "Ouvrir .../app/admin-refusal-reasons.html",
            "Créer un motif : code = TEST_MOTIF , label = « Motif test séance 4 ».",
            "Modifier le label, puis supprimer le motif de test.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "6. Scénario E — Candidat", 1)
    for i, t in enumerate(
        [
            "Se connecter avec candidat@hirelink.ma / Candidat123!",
            "Aller sur Offres.",
            "Ouvrir une offre publiée.",
            "Postuler (suivre le wizard jusqu’à confirmation).",
            "Aller sur Mes candidatures et vérifier le statut Soumise / SUBMITTED.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "7. Scénario F — Postman (tests API Séance 4)", 1)
    for i, t in enumerate(
        [
            "Ouvrir Postman.",
            "File → Import → sélectionner le fichier :",
            "C:\\Users\\daurinia\\Desktop\\HireLink\\docs\\postman\\HireLink_CRUD_Seance4.postman_collection.json",
            "Dans la collection, ouvrir le dossier Auth.",
            "Exécuter « Login recruteur » → Status 200 → le token est enregistré automatiquement.",
            "Exécuter dans l’ordre le dossier « CRUD Offres » : CREATE → READ liste → READ by id → UPDATE → DELETE.",
            "Exécuter « Login admin ».",
            "Exécuter le dossier CRUD Motifs de refus.",
            "Exécuter le dossier CRUD Notifications.",
            "Résultat attendu : toutes les requêtes en 200 (ou 201) sans erreur rouge.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "8. Vérifier les tables (phpMyAdmin) — sans confusion", 1)
    for i, t in enumerate(
        [
            "Ouvrir http://localhost/phpmyadmin",
            "Cliquer sur la base hirelink à gauche.",
            "Vérifier la présence des tables : users, job_offers, applications, refusal_reason_catalog, notifications.",
            "Rappel oral : ces tables ont été générées par Hibernate à partir des @Entity Java, pas par un script .sql manuel comme source.",
        ],
        start=1,
    ):
        add_p(doc, f"{i}. {t}", space_after=2)

    add_heading(doc, "9. Si ça ne marche pas", 1)
    add_table(
        doc,
        ["Symptôme", "Action exacte"],
        [
            ["health ne répond pas", "Relancer le jar ; vérifier que le port 8081 est libre"],
            ["Page blanche / auth échoue", "Vérifier que le front tourne sur 5500 et l’API sur 8081"],
            ["Mot de passe refusé", "Recopier exactement les comptes du tableau §1"],
            ["phpMyAdmin vide", "MySQL XAMPP non démarré, ou mauvaise base (choisir hirelink)"],
            ["Postman 401", "Relancer Login recruteur/admin avant les CRUD"],
        ],
    )

    path = OUT / "02_GUIDE_UTILISATION_ET_COMPTES_HireLink.docx"
    doc.save(path)
    return path


def build_tests(results_path: Path | None):
    doc = Document()
    cover_page(
        doc,
        "DOCUMENT DE TESTS — Séances 3 & 4",
        "Cas de test exécutables · Postman · UI · BDD",
    )

    add_heading(doc, "1. Environnement de test", 1)
    add_table(
        doc,
        ["Élément", "Valeur"],
        [
            ["OS", "Windows"],
            ["API", "http://localhost:8081/api/v1"],
            ["Frontend", "http://localhost:5500"],
            ["BDD", "MySQL/MariaDB · schéma hirelink"],
            ["Collection Postman", "docs/postman/HireLink_CRUD_Seance4.postman_collection.json"],
            ["Date du pack", TODAY],
        ],
    )

    add_heading(doc, "2. Matrice Séance 3", 1)
    add_table(
        doc,
        ["ID", "Cas de test", "Résultat attendu", "Statut"],
        [
            ["S3-01", "Démarrer l’API avec ddl-auto=update", "Tables présentes dans hirelink", "À exécuter / OK si tables visibles"],
            ["S3-02", "Ouvrir JobOffer.java", "Classe @Entity avec champs mappés", "OK code"],
            ["S3-03", "Vérifier relation JobOffer↔Tag", "@ManyToMany + table job_offer_tags", "OK code"],
            ["S3-04", "Vérifier Membership User↔Organization", "Table memberships", "OK code + BDD"],
            ["S3-05", "Diagramme de classes livré", "PNG FINAL présent", "OK"],
        ],
    )

    add_heading(doc, "3. Matrice Séance 4 — API (Postman)", 1)
    add_table(
        doc,
        ["ID", "Cas", "Attendu", "Statut"],
        [
            ["S4-A1", "POST /auth/login recruteur", "200 + accessToken", "voir résultats"],
            ["S4-O1", "POST /offers", "200/201 + id", "voir résultats"],
            ["S4-O2", "GET /offers/mine", "200 + liste", "voir résultats"],
            ["S4-O3", "GET /offers/{id}", "200 + objet", "voir résultats"],
            ["S4-O4", "PUT /offers/{id}", "200 + titre modifié", "voir résultats"],
            ["S4-O5", "DELETE /offers/{id}", "200", "voir résultats"],
            ["S4-A2", "POST /auth/login admin", "200 + accessToken", "voir résultats"],
            ["S4-R1", "CRUD refusal-reasons", "200 sur create/read/update/delete", "voir résultats"],
            ["S4-N1", "CRUD notifications", "200 sur create/read/read/delete", "voir résultats"],
        ],
    )

    if results_path and results_path.exists():
        add_heading(doc, "4. Résultats d’exécution automatisée", 1)
        text = results_path.read_text(encoding="utf-8", errors="replace")
        # keep short
        add_p(doc, text[:4500])

    add_heading(doc, "5. Matrice UI (manuel, suivi professeur)", 1)
    add_table(
        doc,
        ["ID", "Action", "Attendu"],
        [
            ["UI-01", "Login recruteur via auth.html", "Arrivée sur recruiter.html"],
            ["UI-02", "Créer / modifier / supprimer offre", "Liste mise à jour à chaque étape"],
            ["UI-03", "Ouvrir pipeline", "Kanban visible"],
            ["UI-04", "Login admin + CRUD motif", "Motif apparaît puis disparaît"],
            ["UI-05", "Login candidat + postuler", "Candidature visible dans Mes candidatures"],
            ["UI-06", "phpMyAdmin hirelink", "Tables job_offers / applications visibles"],
        ],
    )

    add_heading(doc, "6. Critère de réussite globale", 1)
    add_p(
        doc,
        "Séance 3 OK si : entities + relations + tables MySQL + diagramme.\n"
        "Séance 4 OK si : 3 CRUD API + UI offres + Postman vert + guide utilisable sans aide.",
        bold=True,
    )

    path = OUT / "03_DOCUMENT_TESTS_SEANCE3_SEANCE4_HireLink.docx"
    doc.save(path)
    return path


def build_readme_pack():
    p = OUT / "00_LIRE_EN_PREMIER.txt"
    p.write_text(
        f"""PACK LIVRABLES — HireLink Séances 3 & 4
Date : {TODAY}

Ordre de lecture pour le professeur :
1) 02_GUIDE_UTILISATION_ET_COMPTES_HireLink.docx   ← commencer ICI (comptes + étapes exactes)
2) 03_DOCUMENT_TESTS_SEANCE3_SEANCE4_HireLink.docx ← matrice de tests + résultats
3) 01_RAPPORT_SEANCE3_SEANCE4_HireLink.docx        ← ce qui est demandé / ce qui a été fait

Collection Postman :
../postman/HireLink_CRUD_Seance4.postman_collection.json

Phrase orale critique (tables) :
Les tables sont générées par Hibernate à partir des classes @Entity
(spring.jpa.hibernate.ddl-auto=update). phpMyAdmin sert à vérifier.
""",
        encoding="utf-8",
    )
    return p


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    results = OUT / "postman_results.txt"
    r1 = build_rapport()
    r2 = build_guide()
    r3 = build_tests(results if results.exists() else None)
    r0 = build_readme_pack()
    print("OK", r0)
    print("OK", r1)
    print("OK", r2)
    print("OK", r3)


if __name__ == "__main__":
    main()
