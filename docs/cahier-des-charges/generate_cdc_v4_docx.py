# -*- coding: utf-8 -*-
"""Génère le Cahier des charges HireLink v4 — clair, tabulaire, DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUT_DIRS = [
    Path(r"C:\Users\daurinia\Desktop\HireLink\docs\cahier-des-charges"),
    Path(r"C:\Users\daurinia\Desktop\HireLink-cahier-des-charges"),
    Path(r"C:\Users\daurinia\Desktop"),
]
FILENAME = "HireLink_Cahier_des_charges_v4.docx"

ROYAL = RGBColor(0x0B, 0x1A, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "0B1A4D"
ALT_BG = "F3F5FA"


def set_run(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    sizes = {1: 16, 2: 13, 3: 11}
    set_run(run, bold=True, size=sizes.get(level, 11), color=ROYAL)
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, bold=bold, size=size)
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        run = p.add_run(it)
        set_run(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run(run, bold=bold, size=size, color=color)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, color=WHITE, size=9)
        shade_cell(cell, HEADER_BG)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            set_cell_text(cell, str(val), size=9)
            if r_i % 2 == 1:
                shade_cell(cell, ALT_BG)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return table


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ——— Page de titre ———
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("HireLink")
    set_run(r, bold=True, size=32, color=ROYAL)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Cahier des charges")
    set_run(r, bold=True, size=18, color=ROYAL)

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run("ATS bilatéral — pipeline, confiance, IA ciblée")
    set_run(r, size=12, italic=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(28)
    r = meta.add_run(
        "EL ASSAL Douae\n"
        "ENSA Tétouan — Université Abdelmalek Essaâdi\n"
        "Module Frameworks Technologie Web\n"
        "Année universitaire 2026–2027\n"
        "Version 4.0 — document opérationnel"
    )
    set_run(r, size=11)

    doc.add_page_break()

    # ——— TOC ———
    add_heading_custom(doc, "Table des matières", 1)
    toc = [
        "Manifeste produit",
        "1. Contexte, problématique, objectifs",
        "2. Benchmarking concurrentiel",
        "3. Acteurs, espaces, RBAC",
        "4. Fonctionnalités (tableau maître)",
        "5. Pipeline de candidature",
        "6. Notifications",
        "7. Mécanismes comportementaux (théorie → produit)",
        "8. Doctrine IA et crédits",
        "9. Contrat de Confiance Bilatéral",
        "10. Croissance SaaS (PLG)",
        "11. Exigences non fonctionnelles",
        "12. Architecture et stack",
        "13. Modèle de données (tables + colonnes)",
        "14. Conclusion",
    ]
    add_bullets(doc, toc)
    doc.add_page_break()

    # ——— Manifeste ———
    add_heading_custom(doc, "Manifeste produit", 1)
    add_para(
        doc,
        "HireLink n’est pas « un ATS avec de l’IA en plus ». Trois refus structurent le produit :",
    )
    add_table(
        doc,
        ["Refus", "Règle", "Conséquence produit"],
        [
            [
                "IA décorative",
                "Pas d’IA sans friction mesurable",
                "4 usages IA seulement (matching, profil, grille, langage)",
            ],
            [
                "Kanban = process",
                "Scorecard + checklists + timeline immuable",
                "Transitions contrôlées, pas un tableau de cartes",
            ],
            [
                "Asymétrie de confiance",
                "L’organisation aussi est évaluée",
                "Score + badge de fiabilité publics",
            ],
        ],
        [3.5, 5.5, 6.5],
    )

    # ——— 1 ———
    add_heading_custom(doc, "1. Contexte, problématique, objectifs", 1)
    add_heading_custom(doc, "1.1 Contexte", 2)
    add_para(
        doc,
        "Au Maroc (PME, ETI, stages / PFE / premier emploi), l’acquisition de candidatures "
        "est couverte par les jobboards et LinkedIn. Le pilotage interne (statuts, délais, "
        "collaboration décideurs, audit des refus) reste faible hors grands comptes. "
        "HireLink cible ce vide : un ATS bilatéral léger, MySQL, sans API payante.",
    )
    add_heading_custom(doc, "1.2 Problématique", 2)
    add_bullets(
        doc,
        [
            "Pas de référentiel unique de candidatures → perte de traçabilité.",
            "Candidat sans statut opérationnel fiable (asymétrie d’information).",
            "Rejets non formalisés → non auditables.",
            "KPI RH rarement industrialisés hors grands comptes.",
            "IA marketing en façade → défiance plutôt que valeur.",
        ],
    )
    add_heading_custom(doc, "1.3 Objectifs", 2)
    add_table(
        doc,
        ["ID", "Objectif"],
        [
            ["O1", "Pipeline à états contrôlés"],
            ["O2", "3 espaces : candidat / organisation multi-recruteurs / admin"],
            ["O3", "Engagement de délai (SLA) + indicateur de respect"],
            ["O4", "Motif de rejet structuré obligatoire"],
            ["O5", "Aide à la décision locale, jamais autonome, sans API payante"],
            ["O6", "Architecture sécurisée, déployable, OpenAPI"],
            ["O7", "Gouvernance multi-acteurs + notifications anti-saturation + confiance bilatérale"],
            ["O8", "Chaque mécanisme comportemental / growth a une origine citée (pas d’arbitraire)"],
        ],
        [1.5, 14],
    )

    # ——— 2 Benchmarking ———
    add_heading_custom(doc, "2. Benchmarking concurrentiel", 1)
    add_para(
        doc,
        "Grille exhaustive des acteurs du marché (Maroc + internationaux). "
        "Critères : acquisition, pipeline, transparence candidat, SLA, motif de rejet, "
        "IA, confiance bilatérale, accessibilité PME / stage–PFE.",
    )
    add_table(
        doc,
        [
            "Plateforme",
            "Catégorie",
            "Points forts",
            "Limites vs HireLink",
        ],
        [
            [
                "ReKrute",
                "Jobboard MA",
                "Audience locale, volume",
                "Pas de pipeline ATS, pas de SLA, pas de motif de rejet",
            ],
            [
                "Emploi.ma",
                "Jobboard MA",
                "Visibilité nationale",
                "Diffusion d’annonces ; orchestration interne absente",
            ],
            [
                "Dreamjob.ma",
                "Jobboard MA",
                "Cible étudiants / juniors",
                "Peu de gouvernance de sélection",
            ],
            [
                "Indeed Maroc",
                "Agrégateur / jobboard",
                "Portée large, alertes",
                "Pas d’ATS bilatéral ni engagement de délai",
            ],
            [
                "LinkedIn / Recruiter",
                "Réseau pro",
                "Sourcing qualifié, marque employeur",
                "Coût élevé ; pipeline faible hors offre entreprise",
            ],
            [
                "Bayt.com",
                "Jobboard MENA",
                "Couverture régionale",
                "Pas positionné process ATS léger PME MA",
            ],
            [
                "Anapec / portails publics",
                "Emploi public",
                "Cadre institutionnel",
                "Hors scope ATS privé PME",
            ],
            [
                "Greenhouse",
                "ATS US",
                "Scorecard, structured hiring",
                "Complexité / coût ; pas confiance bilatérale publique",
            ],
            [
                "Lever",
                "ATS US",
                "CRM candidat, kanban fluide",
                "Peu adapté MENA / stage–PFE ; asymétrie confiance",
            ],
            [
                "Teamtailor",
                "ATS + career site",
                "Pages carrière, UX candidat",
                "Pas de SLA contractualisé ni réciprocité de score",
            ],
            [
                "Workable",
                "ATS SaaS",
                "Équilibre features / prix",
                "IA généraliste ; pas de badge fiabilité org",
            ],
            [
                "SmartRecruiters",
                "ATS entreprise",
                "Marketplace, reporting",
                "Cible grands comptes",
            ],
            [
                "BambooHR (recrutement)",
                "HRIS + ATS léger",
                "Intégration RH",
                "Module recrutement peu profond",
            ],
            [
                "Ashby / Pinpoint",
                "ATS modernes",
                "Analytics, UX soignée",
                "Prix / complexité hors PME–junior MA",
            ],
            [
                "HRMAPS",
                "ATS Maroc",
                "Ancrage local",
                "Orchestration / preuves / SLA publics limités",
            ],
            [
                "CBO / solutions locales",
                "ATS / RH MA",
                "Présence marché",
                "Peu de science comportementale citée ; IA opaque",
            ],
            [
                "HireLink (cible)",
                "ATS bilatéral",
                "SLA, motif rejet, IA locale à crédits, score confiance 2 voies, PLG",
                "—",
            ],
        ],
        [3.2, 2.8, 4.5, 5.0],
    )

    add_heading_custom(doc, "2.1 Espace blanc retenu", 2)
    add_bullets(
        doc,
        [
            "SLA visible du candidat + indice de réactivité organisation.",
            "Confiance bilatérale (score org public + score candidat non punitif).",
            "IA locale dosée (crédits), 4 usages seulement.",
            "Organisation multi-recruteurs pour structures légères.",
        ],
    )

    # ——— 3 Acteurs ———
    add_heading_custom(doc, "3. Acteurs, espaces, RBAC", 1)
    add_para(
        doc,
        "Entité racine = Organisation (workspace), pas un compte recruteur isolé.",
    )
    add_table(
        doc,
        ["Espace", "Contenu"],
        [
            [
                "Candidat",
                "Profil, offres, candidatures, notifs, NPS, score fiabilité, parrainage",
            ],
            [
                "Organisation",
                "Owner, Recruteur, Hiring Manager, Panéliste invité (lien limité)",
            ],
            [
                "Admin plateforme",
                "Vérif organisations, modération, quotas IA globaux — sans croisement nominatif",
            ],
        ],
        [3.5, 12],
    )
    add_heading_custom(doc, "3.1 Matrice RBAC", 2)
    add_table(
        doc,
        ["Rôle", "Peut", "Ne peut pas"],
        [
            [
                "Owner",
                "Inviter/retirer, facturation, crédits, offres, tout voir",
                "—",
            ],
            [
                "Recruteur",
                "Offres, pipeline, entretiens, offres salariales",
                "Facturation / gestion membres",
            ],
            [
                "Hiring Manager",
                "Scorecard, entretiens métier, commentaires, checklist assignée",
                "Créer offre seul / gérer membres",
            ],
            [
                "Panéliste",
                "1 dossier + note scorecard",
                "Voir les autres candidatures",
            ],
            [
                "Candidat",
                "Profil, postuler, suivre, export RGPD",
                "Voir données d’autres candidats / pipeline interne",
            ],
        ],
        [3, 7, 5.5],
    )

    # ——— 4 Fonctionnalités ———
    add_heading_custom(doc, "4. Fonctionnalités (tableau maître)", 1)
    add_para(doc, "Liste précise livrable. Priorité : H = haute, M = moyenne.")
    add_table(
        doc,
        ["ID", "Fonctionnalité", "Acteur", "Règle / méthode", "Prio"],
        [
            ["F01", "Auth email + OTP + JWT", "Tous", "OTP à usage unique, JWT signé", "H"],
            ["F02", "OAuth Google (option)", "Tous", "Token vérifié serveur", "M"],
            ["F03", "Profil candidat + OCR CV local", "Candidat", "Extraction skills locale", "H"],
            ["F04", "Complétion profil %", "Candidat", "Zeigarnik / gradient d’objectif", "H"],
            ["F05", "Publication offre + SLA jours", "Recruteur", "SLA stocké sur l’offre", "H"],
            ["F06", "Détecteur formulation genrée", "Recruteur", "Gaucher et al. 2011 — lexique local", "H"],
            ["F07", "Matching candidat↔offre", "Les deux", "Jaccard / couverture — tri humain", "H"],
            ["F08", "Candidature multi-étapes", "Candidat", "CV → lettre opt. → envoi", "H"],
            ["F09", "Pipeline états contrôlés", "Org", "Transitions serveur uniquement", "H"],
            ["F10", "Checklists bloquantes", "Org", "Pas d’avance si items obligatoires KO", "H"],
            ["F11", "Scorecard pondérée", "Org / HM", "Entretien structuré (Levashina 2014)", "H"],
            ["F12", "Multi-entretiens typés", "Org", "PHONE / VISIO / ONSITE / ASYNC", "H"],
            ["F13", "Timeline immuable", "Org / Candidat", "Append-only des transitions", "H"],
            ["F14", "Rejet = motif catalogue", "Org", "Justice procédurale (Colquitt)", "H"],
            ["F15", "SLA visible + retard", "Les deux", "Aversion à la perte (cadrage)", "H"],
            ["F16", "Rappels si-alors", "Org / Candidat", "Intentions Gollwitzer 1999", "H"],
            ["F17", "Inbox prioritaire SLA·score", "Recruteur", "Tri serveur", "H"],
            ["F18", "Dashboard petites victoires", "Recruteur", "Progress Principle (Amabile)", "M"],
            ["F19", "Suggestions profil IA", "Candidat", "Suggestions only, pas d’auto-write", "H"],
            ["F20", "Grille entretien IA", "Recruteur", "Banque locale + crédits", "H"],
            ["F21", "Ledger crédits Free/Premium", "Les deux", "Effet de dotation (Thaler)", "H"],
            ["F22", "Score fiabilité org", "Org / public", "SLA + décisions explicites", "H"],
            ["F23", "Score fiabilité candidat", "Candidat", "Non bloquant ; silence ≠ annulation", "H"],
            ["F24", "Badge fiabilité embarquable", "Org", "PLG / backlink Trustpilot-like", "M"],
            ["F25", "Parrainage candidat→candidat", "Candidat", "Bonus après profil complété", "M"],
            ["F26", "Onboarding checklist org", "Owner", "Zeigarnik activation", "M"],
            ["F27", "Notifs in-app + email + digest", "Tous", "Charge cognitive (Sweller)", "H"],
            ["F28", "Préférences notifs", "Tous", "Contractuelles non désactivables", "H"],
            ["F29", "Export / consentements RGPD", "Candidat", "Loi 09-08 / autonomie", "H"],
            ["F30", "Vérif entreprise (ICE)", "Org / Admin", "Statut vérifié affiché", "H"],
            ["F31", "Admin quotas / modération", "Admin", "Isolation multi-tenant", "M"],
            ["F32", "Blacklist soft", "Org", "Non punitive, auditable", "M"],
            ["F33", "Offre salariale tracée", "Org / Candidat", "Peak-End Rule (clôture claire)", "M"],
            ["F34", "NPS post-entretien", "Candidat", "Feedback structuré", "M"],
            ["F35", "Paywall contextuel crédits", "Les deux", "Upgrade au moment de friction", "M"],
        ],
        [1.2, 4.2, 2.2, 5.5, 1.2],
    )

    # ——— 5 Pipeline ———
    add_heading_custom(doc, "5. Pipeline de candidature", 1)
    add_table(
        doc,
        ["État", "Signification", "Transition autorisée vers", "Contrôle serveur"],
        [
            ["SUBMITTED", "Dossier déposé", "UNDER_REVIEW, WITHDRAWN", "Auth candidat / org"],
            ["UNDER_REVIEW", "Présélection", "INTERVIEW_SCHEDULED, REJECTED, WITHDRAWN", "Checklist + rôle"],
            ["INTERVIEW_SCHEDULED", "Entretien planifié", "INTERVIEW_COMPLETED, REJECTED, WITHDRAWN", "Créneau + type"],
            ["INTERVIEW_COMPLETED", "Entretien tenu", "ACCEPTED, REJECTED, INTERVIEW_SCHEDULED", "Scorecard min."],
            ["ACCEPTED", "Décision positive", "— (terminal)", "Motif / offre salariale opt."],
            ["REJECTED", "Décision négative", "— (terminal)", "Motif catalogue obligatoire"],
            ["WITHDRAWN", "Retrait candidat", "— (terminal)", "Candidat uniquement"],
        ],
        [3.5, 3.5, 5, 3.5],
    )
    add_bullets(
        doc,
        [
            "Aucune transition depuis le front seul : l’API valide rôle, ownership, checklist, motif.",
            "Timeline append-only : chaque transition = événement (alimente notifs + fiabilité).",
            "SLA : deadline = submitted_at + offer.response_sla_days ; dépassement visible candidat.",
        ],
    )

    # ——— 6 Notifs ———
    add_heading_custom(doc, "6. Notifications", 1)
    add_table(
        doc,
        ["Canal", "Usage", "Règle"],
        [
            ["In-app", "Tous événements pertinents", "Toujours actif, historisé"],
            ["Email", "Statut, entretien, offre, rappels J-1/J-2", "Désactivable sauf contractuel"],
            ["Push web", "Temps réel optionnel", "Opt-in"],
            ["WhatsApp", "Rappel entretien", "Consentement distinct"],
            ["Digest quotidien", "Volume candidat→recruteur", "Immédiat seulement si dossier suivi"],
        ],
        [3.5, 6, 6],
    )

    # ——— 7 Méthodes (compact table ONLY) ———
    add_heading_custom(doc, "7. Mécanismes comportementaux (théorie → produit)", 1)
    add_para(
        doc,
        "Pas de chapitre théorique autonome : chaque méthode est citée ici une fois, "
        "mappée à un mécanisme produit. Garde-fous : réversibilité, transparence, alignement d’intérêt "
        "(pas de dark pattern).",
    )
    add_table(
        doc,
        ["Méthode / origine", "Mécanisme HireLink", "Effet attendu"],
        [
            [
                "Justice procédurale (Leventhal ; Colquitt)",
                "Motif de rejet obligatoire + critères visibles",
                "Refus perçu plus équitable",
            ],
            [
                "Entretiens structurés (Levashina et al.)",
                "Scorecard + grille d’entretien",
                "Moins de halo / meilleure prédictivité",
            ],
            [
                "Peak-End Rule (Kahneman)",
                "Clôture structurée (feedback / offre claire)",
                "Meilleure perception globale du process",
            ],
            [
                "Aversion à la perte (Kahneman & Tversky)",
                "SLA cadré comme rupture possible",
                "Plus d’action côté recruteur",
            ],
            [
                "Zeigarnik + gradient d’objectif",
                "% profil + checklists incomplètes",
                "Complétion motivée",
            ],
            [
                "Intentions si-alors (Gollwitzer)",
                "Rappels « Si … alors … »",
                "Meilleur taux d’exécution",
            ],
            [
                "Autodétermination (Deci & Ryan)",
                "IA = suggestion ; export RGPD 1 clic",
                "Autonomie préservée",
            ],
            [
                "Formulation genrée (Gaucher et al. 2011)",
                "Language-check avant publication",
                "Bassin de candidatures élargi",
            ],
            [
                "Progress Principle (Amabile & Kramer)",
                "Dashboard « petites victoires »",
                "Motivation / rétention RH",
            ],
            [
                "Preuve sociale (Cialdini)",
                "Score fiabilité org public",
                "Pression d’alignement",
            ],
            [
                "Effet de dotation (Thaler)",
                "Crédits Free = capital déjà possédé",
                "Attachement palier gratuit",
            ],
            [
                "Charge cognitive (Sweller)",
                "Digest vs immédiat",
                "Attention préservée",
            ],
        ],
        [5, 5.5, 5],
    )

    # ——— 8 IA ———
    add_heading_custom(doc, "8. Doctrine IA et crédits", 1)
    add_para(doc, "Une capacité IA n’entre que si les 3 conditions sont vraies ensemble.")
    add_table(
        doc,
        ["Condition", "Test", "Échoue si…"],
        [
            ["1. Friction réelle", "Temps / qualité perdus sans aide ?", "Chatbot sans tâche"],
            ["2. Méthode locale bornée", "Règles / similarité sans LLM payant ?", "Texte libre halluciné"],
            ["3. Humain décide", "IA trie / suggère, ne décide pas ?", "Auto-rejet sur score"],
        ],
        [4, 6, 5.5],
    )
    add_heading_custom(doc, "8.1 Quatre usages retenus", 2)
    add_table(
        doc,
        ["Usage", "Méthode locale", "Humain"],
        [
            ["Matching profil ↔ offre", "Jaccard / couverture compétences", "Tri, jamais filtre auto"],
            ["Suggestions de profil", "Écart vs exigences marché", "Candidat accepte / ignore"],
            ["Grille d’entretien", "Banque de questions locale", "Manager édite et mène"],
            ["Language-check offre", "Marqueurs lexicaux connotés", "Recruteur publie ou corrige"],
        ],
        [4.5, 6, 5],
    )
    add_heading_custom(doc, "8.2 Exclus", 2)
    add_bullets(
        doc,
        [
            "Chatbot généraliste.",
            "Rédaction libre de refus (hallucination + justice procédurale).",
            "Score « embaucher X » automatisé.",
        ],
    )
    add_heading_custom(doc, "8.3 Crédits", 2)
    add_table(
        doc,
        ["Palier", "Candidat", "Organisation"],
        [
            ["Free", "5 crédits / mois", "20 crédits / mois (partagés)"],
            ["Premium", "Illimité + suggestions approfondies", "Illimité + grilles sur pipeline actif"],
        ],
        [3, 6, 6.5],
    )

    # ——— 9 Confiance ———
    add_heading_custom(doc, "9. Contrat de Confiance Bilatéral", 1)
    add_para(
        doc,
        "Score bidirectionnel calculé depuis la timeline — jamais une note manuelle subjective.",
    )
    add_table(
        doc,
        ["Côté", "Facteurs", "Garde-fous"],
        [
            [
                "Organisation",
                "Taux SLA tenus + taux de décisions explicites",
                "Public ; badge optionnel",
            ],
            [
                "Candidat",
                "Présence entretiens ; réponse offre ; retrait vs silence",
                "Non bloquant ; annulation anticipée non pénalisée",
            ],
        ],
        [3, 7, 5.5],
    )
    add_para(
        doc,
        "Badge HTML embarquable (page carrière) : marque employeur + backlink HireLink (PLG).",
    )

    # ——— 10 Growth ———
    add_heading_custom(doc, "10. Croissance SaaS (PLG)", 1)
    add_table(
        doc,
        ["Boucle", "Mécanique", "Déclencheur"],
        [
            ["Activation", "1er score matching / 1re checklist template", "Time-to-value court"],
            ["Expansion quota", "Compteur crédits + paywall au moment du manque", "Dotation + usage-based"],
            ["Parrainage", "Bonus crédit après profil invité complété", "Anti-spam inscription"],
            ["Badge viral", "Snippet org sur site carrière", "Backlink / preuve sociale"],
            ["Onboarding checklist", "Créer offre → inviter → publier", "Zeigarnik activation"],
            ["Rétention", "Digest progrès avant backlog", "Petites victoires"],
        ],
        [3.5, 6.5, 5.5],
    )
    add_para(
        doc,
        "Pricing : socle ATS illimité en Free ; monétisation = couche IA (usage), pas le siège.",
    )

    # ——— 11 NFR ———
    add_heading_custom(doc, "11. Exigences non fonctionnelles", 1)
    add_table(
        doc,
        ["Famille", "Exigence"],
        [
            ["Sécurité", "BCrypt/Argon2 ; JWT ; RBAC ; validation entrées ; isolation CV ; anti-IDOR pipeline"],
            ["Confidentialité", "Consentements ; finalités limitées ; isolation organisations"],
            ["Fiabilité", "Transactions sur transitions ; migrations reproductibles (Flyway)"],
            ["Observabilité", "Logs structurés ; corrélation requêtes"],
            ["Documentation", "OpenAPI 3 synchronisée"],
            ["Performance", "Latences nominales acceptables parcours critiques"],
            ["Équité algo", "Score fiabilité recalculable, aucun facteur caché"],
            ["Notifs", "Plafond d’immédiats hors dossiers suivis"],
            ["Non-manipulation", "Garde-fous §7 vérifiables en revue avant prod"],
        ],
        [3.5, 12],
    )

    # ——— 12 Archi ———
    add_heading_custom(doc, "12. Architecture et stack", 1)
    add_table(
        doc,
        ["Composant", "Choix"],
        [
            ["Langage", "Java 21 LTS"],
            ["Framework", "Spring Boot 3.x"],
            ["API", "REST + OpenAPI (springdoc)"],
            ["Sécurité", "Spring Security 6 + JWT stateless"],
            ["Persistance", "Spring Data JPA / Hibernate 6"],
            ["SGBD", "MySQL 8.4"],
            ["Migrations", "Flyway"],
            ["Build / tests", "Maven ; JUnit 5 ; MockMvc / Testcontainers"],
            ["Conteneurs", "Docker Compose"],
            ["IA", "Module Java local — aucune API payante"],
        ],
        [4, 11.5],
    )
    add_para(
        doc,
        "Modules : notification · ai-copilot (+ credit-ledger) · reliability · growth · pipeline / auth.",
    )

    # ——— 13 Data model DETAILED ———
    add_heading_custom(doc, "13. Modèle de données", 1)
    add_para(
        doc,
        "Chaque table : utilité métier + colonnes clés. SGBD cible MySQL 8.",
    )

    tables_spec = [
        (
            "users",
            "Comptes authentifiés (candidat, membres org, admin).",
            [
                ["id", "PK"],
                ["email", "Unique, login"],
                ["password_hash", "BCrypt"],
                ["full_name", "Affichage"],
                ["role", "Enum applicatif de base"],
                ["company / company_ice", "Org recruteur"],
                ["company_verified", "Bool vérif admin"],
                ["skills, sought_roles, …", "Profil candidat"],
                ["created_at", "Audit"],
            ],
        ),
        (
            "organizations",
            "Workspace entreprise (racine multi-membres).",
            [
                ["id", "PK"],
                ["name", "Raison sociale"],
                ["domain / ice", "Vérification"],
                ["owner_user_id", "FK users"],
                ["reliability_score_cache", "Dénormalisation optionnelle"],
            ],
        ),
        (
            "memberships",
            "RBAC org : lien user ↔ organisation + rôle.",
            [
                ["id", "PK"],
                ["organization_id", "FK"],
                ["user_id", "FK"],
                ["org_role", "OWNER / RECRUITER / HM"],
                ["invited_at / accepted_at", "Parcours invitation"],
            ],
        ),
        (
            "panelist_access_links",
            "Accès ponctuel 1 dossier sans compte complet.",
            [
                ["token", "Secret URL"],
                ["application_id", "FK dossier"],
                ["expires_at", "TTL"],
                ["scopes", "READ + SCORECARD_SUBMIT"],
            ],
        ),
        (
            "job_offers",
            "Offres publiées avec engagement de délai.",
            [
                ["id", "PK"],
                ["organization_id / recruiter_id", "FK"],
                ["title, description, skills", "Contenu"],
                ["response_sla_days", "Engagement délai"],
                ["published", "Visibilité"],
                ["created_at", "Audit"],
            ],
        ),
        (
            "applications",
            "Candidature = objet central du pipeline.",
            [
                ["id", "PK"],
                ["candidate_id", "FK users"],
                ["job_offer_id", "FK"],
                ["status", "État machine"],
                ["match_score", "Aide tri (nullable)"],
                ["submitted_at / updated_at", "SLA + audit"],
            ],
        ),
        (
            "application_status_history",
            "Timeline immuable (append-only).",
            [
                ["id", "PK"],
                ["application_id", "FK"],
                ["from_status / to_status", "Transition"],
                ["actor_user_id", "Qui"],
                ["payload_json", "Motif, checklist proof"],
                ["created_at", "Horodatage"],
            ],
        ),
        (
            "interviews",
            "Entretiens multi-types.",
            [
                ["id", "PK"],
                ["application_id", "FK"],
                ["type", "PHONE/VISIO/ONSITE/ASYNC"],
                ["scheduled_at / location", "Logistique"],
                ["panelist_ids", "Participants"],
            ],
        ),
        (
            "refusal_reason_catalog / application_refusals",
            "Motifs normalisés obligatoires au rejet.",
            [
                ["code / label", "Catalogue"],
                ["application_id", "FK refus"],
                ["reason_code", "FK catalogue"],
                ["comment", "Optionnel borné"],
            ],
        ),
        (
            "notification_preferences / notification_deliveries",
            "Canaux + journal d’envoi / digest.",
            [
                ["user_id", "FK"],
                ["channel / enabled", "Préférences"],
                ["event_type / sent_at / digest_bucket", "Livraison"],
            ],
        ),
        (
            "ai_credit_ledgers",
            "Compteur mensuel Free / Premium.",
            [
                ["user_or_org_id", "Sujet"],
                ["period_ym", "Mois"],
                ["used / limit", "Quota"],
                ["action", "Type consommation"],
            ],
        ),
        (
            "reliability_events",
            "Événements alimentant les scores (dérivés timeline).",
            [
                ["subject_type / subject_id", "ORG ou CANDIDATE"],
                ["event_type", "SLA_OK, DECISION, NO_SHOW…"],
                ["weight / created_at", "Calcul auditable"],
            ],
        ),
        (
            "reliability_badges",
            "Badge public embarquable.",
            [
                ["organization_id", "FK"],
                ["token", "Intégration externe"],
                ["score_snapshot", "Affiché"],
                ["view_count", "PLG métrique"],
            ],
        ),
        (
            "referral_links",
            "Parrainage candidat → candidat.",
            [
                ["code", "URL ?ref="],
                ["referrer_user_id", "FK"],
                ["status", "INVITED / CONVERTED"],
                ["bonus_granted", "Crédit IA"],
            ],
        ),
        (
            "job_offer_language_checks",
            "Résultat language-check avant publish.",
            [
                ["job_offer_id", "FK"],
                ["markers_json / suggestions_json", "Sortie"],
                ["accepted", "Bool recruteur"],
                ["checked_at", "Horodatage"],
            ],
        ),
        (
            "matching_analyses / consent_records / engagement_events",
            "Matching auditable ; consentements ; events growth (profil %, onboarding).",
            [
                ["application_or_user_id", "Sujet"],
                ["payload_json", "Détail score / consent / step"],
                ["created_at", "Audit"],
            ],
        ),
    ]

    for name, utility, cols in tables_spec:
        add_heading_custom(doc, f"Table `{name}`", 2)
        add_para(doc, f"Utilité : {utility}", size=10)
        add_table(doc, ["Colonne / champ", "Rôle"], cols, [6, 9.5])

    # ——— 14 Conclusion ———
    add_heading_custom(doc, "14. Conclusion", 1)
    add_para(
        doc,
        "HireLink = ATS bilatéral (pipeline + SLA + motif de rejet) + IA locale bornée "
        "(4 usages, crédits) + Confiance Bilatérale (score / badge) + PLG (parrainage, onboarding, "
        "petites victoires). Les méthodes scientifiques sont citées dans le tableau §7 — "
        "elles justifient des mécanismes produit, elles ne forment pas un cours théorique. "
        "Stack imposée : Spring Boot 3 / MySQL, sans API payante.",
    )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(24)
    r = footer.add_run("— Fin du cahier des charges HireLink v4.0 —")
    set_run(r, italic=True, size=10, color=ROYAL)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paths = []
    for d in OUT_DIRS:
        d.mkdir(parents=True, exist_ok=True)
        p = d / FILENAME
        doc.save(p)
        paths.append(p)
    return paths


if __name__ == "__main__":
    for p in build():
        print("OK", p, p.stat().st_size)
