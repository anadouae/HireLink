# -*- coding: utf-8 -*-
"""
Cahier des charges HireLink — version claire & convaincante
Structure prof + résumé fluide + BDD complète + sources marché.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUT = [
    Path(r"C:\Users\daurinia\Desktop\HireLink\docs\cahier-des-charges"),
    Path(r"C:\Users\daurinia\Desktop"),
    Path(r"C:\Users\daurinia\Downloads"),
]
NAME = "HireLink_Cahier_des_charges.docx"

ROYAL = RGBColor(0x0B, 0x1A, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = "0B1A4D"
ALT = "F3F5FA"
GOLD = RGBColor(0x8A, 0x6A, 0x1F)


def rs(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def h(doc, text, level=1):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    para.paragraph_format.space_after = Pt(4)
    r = para.add_run(text)
    rs(r, bold=True, size={1: 15, 2: 12, 3: 11}[level], color=ROYAL)


def p(doc, text, *, size=11, after=6, bold=False, italic=False, color=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.space_before = Pt(0)
    r = para.add_run(text)
    rs(r, bold=bold, size=size, italic=italic, color=color)
    return para


def bullets(doc, items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        r = para.add_run(it)
        rs(r, size=11)
        para.paragraph_format.space_after = Pt(2)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>'))


def cell_txt(cell, text, *, bold=False, color=None, size=8):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(1)
    para.paragraph_format.space_after = Pt(1)
    r = para.add_run(str(text))
    rs(r, bold=bold, size=size, color=color)


def table(doc, headers, rows, widths=None, size=8):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        cell_txt(c, hd, bold=True, color=WHITE, size=size)
        shade(c, HDR)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            cell_txt(c, val, size=size)
            if ri % 2:
                shade(c, ALT)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return t


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.6)
        s.bottom_margin = Cm(1.6)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)

    # ——— TITRE ———
    for _ in range(2):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs(t.add_run("HireLink"), bold=True, size=28, color=ROYAL)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs(st.add_run("Cahier des charges"), bold=True, size=16, color=ROYAL)

    st2 = doc.add_paragraph()
    st2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs(
        st2.add_run(
            "ATS bilatéral — le recrutement où chaque partie s’engage\n"
            "(délai visible · statut fiable · rejet motivé · décision humaine)"
        ),
        size=11,
        italic=True,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(20)
    rs(
        meta.add_run(
            "EL ASSAL Douae\n"
            "ENSA Tétouan — Université Abdelmalek Essaâdi\n"
            "Module Frameworks Technologie Web · Spring Boot / MySQL\n"
            "Année universitaire 2026–2027 · Version 4.1"
        ),
        size=11,
    )

    doc.add_page_break()

    # ——— RESUME RIVIERE ———
    h(doc, "Résumé — le fil conducteur")
    p(
        doc,
        "Au Maroc, trouver des candidats n’est plus le vrai problème : LinkedIn dépasse "
        "7 millions d’utilisateurs locaux, et les jobboards (ReKrute, Emploi.ma, Dreamjob, Indeed…) "
        "apportent déjà le volume (AfricaWork, classements 2025). Le vrai goulot, c’est ce qui "
        "se passe après le dépôt du CV : dossiers dispersés dans des boîtes mail, statut opaque "
        "pour le candidat, silence prolongé, refus sans explication.",
    )
    p(
        doc,
        "Ce silence a un coût mesurable. Selon SHRM, environ 60 % des chercheurs d’emploi "
        "déclarent avoir été « ghostés » par un employeur ; 72 % en parlent autour d’eux. "
        "LinkedIn (Global Talent Trends) indique que 75 % des candidats jugent la culture "
        "d’une entreprise à travers son process de recrutement. Autrement dit : mal traiter "
        "une candidature dégrade la marque employeur autant qu’un mauvais produit.",
    )
    p(
        doc,
        "Les ATS internationaux (Greenhouse, Lever, Workday…) industrialisent ce process — "
        "mais ils sont chers, lourds, et pensés pour des équipes RH structurées. Les PME, "
        "startups et structures qui recrutent stagiaires / PFE / premiers emplois restent "
        "souvent entre Excel et la messagerie. C’est exactement l’espace que HireLink occupe.",
    )
    p(
        doc,
        "HireLink n’est pas un énième site d’annonces. C’est une plateforme bilatérale : "
        "le même dossier est vu des deux côtés. L’organisation s’engage sur un délai de réponse "
        "(SLA) visible. Chaque rejet exige un motif clair. Chaque transition est tracée. "
        "Un score d’aide au tri (matching local) accélère la lecture — sans jamais décider "
        "à la place du recruteur. Aucune API d’IA payante : contrainte pédagogique respectée, "
        "et CV qui ne sortent pas vers un cloud tiers.",
    )
    p(
        doc,
        "Ce document suit le canevas demandé : domaine → problème & objectifs → utilisateurs → "
        "acteurs → fonctionnalités → règles métier → réflexion → classes → conception de la "
        "base de données. L’annexe positionne HireLink face au marché.",
        italic=True,
        size=10,
    )

    # Atouts box as table
    h(doc, "Ce que HireLink apporte (valeur)", 2)
    table(
        doc,
        ["Atout", "Ce que ça change concrètement"],
        [
            [
                "Pipeline à états contrôlés",
                "Plus de dossier « perdu » : chaque candidature a un statut unique et une histoire",
            ],
            [
                "SLA de réponse visible",
                "L’entreprise s’engage ; le candidat voit le compte à rebours ; le retard est mesuré",
            ],
            [
                "Rejet motivé obligatoire",
                "Fini le silence : tout refus laisse une trace compréhensible et auditable",
            ],
            [
                "Vue bilatérale",
                "Candidat et recruteur regardent le même process — réduction de l’asymétrie",
            ],
            [
                "Matching local (Jaccard)",
                "Aide au tri des CV sans boîte noire cloud ni décision automatique",
            ],
            [
                "Sécurité serveur",
                "Impossible d’avancer un dossier « en trichant » côté navigateur : l’API décide",
            ],
            [
                "Conformité 09-08",
                "Consentements, finalité recrutement, droits d’accès — CVs traités localement",
            ],
        ],
        [4.5, 11],
        size=9,
    )

    doc.add_page_break()

    # TOC
    h(doc, "Table des matières")
    bullets(
        doc,
        [
            "1. Domaine",
            "2. Problématique et objectifs",
            "3. Utilisateurs concernés",
            "4. Identification des acteurs",
            "5. Identification des fonctionnalités",
            "6. Règles de métier",
            "7. Première réflexion sur le domaine",
            "8. Principales classes",
            "9. Conception de la base de données",
            "Annexe A — Benchmarking marché",
            "Annexe B — Sources",
        ],
    )

    # ========== 1 ==========
    h(doc, "1. Domaine")
    p(
        doc,
        "Le domaine d’HireLink est le recrutement digitalisé dans les organisations marocaines "
        "(PME, ETI, cabinets, structures accueillant stages et PFE) : de la publication du besoin "
        "jusqu’à la décision finale, en passant par la réception, la présélection et les entretiens.",
    )
    p(
        doc,
        "On distingue clairement deux familles d’outils déjà présentes sur le marché :",
    )
    bullets(
        doc,
        [
            "Jobboards / réseaux (ReKrute, Emploi.ma, Dreamjob, Indeed, LinkedIn) → apportent du volume.",
            "ATS (Applicant Tracking System) → orchestrent le process interne de sélection.",
        ],
    )
    p(
        doc,
        "HireLink appartient à la seconde famille, avec une exigence supplémentaire : le process "
        "est bilatéral. Le candidat n’est pas un dossier « opaque » ; l’organisation n’est pas "
        "un acteur qui peut disparaître sans laisser de trace.",
    )
    table(
        doc,
        ["Concept métier", "Définition", "Pourquoi c’est central"],
        [
            ["Organisation", "Entreprise qui recrute (espace multi-membres)", "Racine des offres et des droits"],
            ["Offre", "Poste publié + compétences + délai de réponse", "Porte le SLA"],
            ["Candidat", "Personne qui postule et suit ses dossiers", "Deuxième face du process"],
            ["Candidature", "Lien candidat ↔ offre + état courant", "Objet central du système"],
            ["Pipeline", "Suite d’états jusqu’à décision / retrait", "Garantit la traçabilité"],
            ["Entretien", "Échange planifié (visio, présentiel…)", "Étape critique du cycle"],
            ["Décision", "Acceptation ou rejet motivé", "Clôture auditable"],
        ],
        [3.2, 6.5, 5.5],
        size=9,
    )

    # ========== 2 ==========
    h(doc, "2. Problématique et objectifs")
    h(doc, "2.1 Problématique (pourquoi ce produit existe)", 2)
    p(
        doc,
        "Le marché marocain sait attirer des candidatures. AfricaWork (2025) confirme la "
        "densité des canaux : Emploi.ma, ReKrute, LinkedIn Talent, Dreamjob, Indeed, agrégateurs… "
        "Les grands groupes industrialisent déjà avec des ATS (Workday, Taleo, Greenhouse, "
        "Teamtailor…). En revanche, dès qu’on sort des grands comptes, le diagnostic terrain "
        "reste le même :",
    )
    table(
        doc,
        ["Constat terrain", "Preuve / lecture marché", "Effet concret"],
        [
            [
                "Candidatures éclatées (mail, WhatsApp, Excel)",
                "Sans ATS, pas de référentiel unique",
                "Dossiers perdus, double traitement, oubli",
            ],
            [
                "Candidat sans statut fiable",
                "Asymétrie d’information classique ATS vs jobboard",
                "Anxiété, désengagement, mauvaise image",
            ],
            [
                "Silence / ghosting employeur",
                "SHRM : ~60 % des candidats ghostés ; 72 % en parlent",
                "Marque employeur abîmée, vivier qui se tarit",
            ],
            [
                "Process jugé = culture jugée",
                "LinkedIn : 75 % jugent l’entreprise via le recrutement",
                "Perte de talents avant même l’embauche",
            ],
            [
                "ATS lourds / chers",
                "Solutions US/EU pensées grands comptes",
                "PME et stage–PFE restent sous-équipés",
            ],
            [
                "Données candidats sensibles",
                "Loi 09-08 / CNDP : information, finalité, droits",
                "Risque légal si CV envoyés à des IA cloud opaques",
            ],
        ],
        [4, 5.5, 5.5],
        size=8,
    )
    p(
        doc,
        "Question de conception : comment réaliser un ATS léger, bilatéral, sécurisé "
        "(Spring Boot + MySQL, sans API payante), qui contractualise un délai de réponse, "
        "rend le statut visible, impose un motif de rejet, et aide au tri sans remplacer "
        "le jugement humain ?",
        bold=True,
        size=10,
    )

    h(doc, "2.2 Objectifs", 2)
    table(
        doc,
        ["ID", "Objectif", "Critère de succès (démo)"],
        [
            ["O1", "Pipeline à états contrôlés", "On ne peut pas sauter une étape illégale"],
            ["O2", "Espaces candidat / organisation / admin", "3 parcours distincts utilisables"],
            ["O3", "SLA par offre, visible des deux côtés", "Retard calculé et affiché"],
            ["O4", "Rejet = motif catalogue obligatoire", "Impossible de rejeter « en silence »"],
            ["O5", "Matching local d’aide au tri", "Score affiché ; jamais auto-rejet"],
            ["O6", "Sécurité JWT + isolation org", "IDOR pipeline bloqué"],
            ["O7", "Traçabilité + conformité de base 09-08", "Timeline + consentements"],
        ],
        [1.2, 6.5, 7.5],
        size=9,
    )

    # ========== 3 ==========
    h(doc, "3. Utilisateurs concernés")
    p(doc, "Qui se connecte réellement à HireLink, et qu’est-ce qu’il vient chercher ?")
    table(
        doc,
        ["Utilisateur", "Profil typique", "Il vient pour…", "Il quitte si…"],
        [
            [
                "Candidat",
                "Étudiant, junior, salarié en mobilité",
                "Postuler, suivre, recevoir une vraie réponse",
                "Silence, statut flou, process opaque",
            ],
            [
                "Recruteur / RH",
                "PME, startup, service RH léger",
                "Ne plus perdre de dossiers, tenir les délais",
                "Outil trop lourd ou non fiable",
            ],
            [
                "Hiring manager",
                "Manager métier invité sur un poste",
                "Noter / commenter sans gérer toute l’offre",
                "Trop de friction pour une intervention ponctuelle",
            ],
            [
                "Administrateur",
                "Ops plateforme / encadrant projet",
                "Vérifier orgs, voir la santé globale",
                "Accès croisé dangereux aux dossiers privés",
            ],
        ],
        [3, 4, 4.5, 4],
        size=8,
    )

    # ========== 4 ==========
    h(doc, "4. Identification des acteurs")
    p(
        doc,
        "L’acteur est le rôle vu du système (plusieurs personnes peuvent partager le même rôle).",
    )
    table(
        doc,
        ["Acteur", "Type", "Ce qu’il peut faire", "Ce qu’il ne peut pas faire"],
        [
            ["Candidat", "Principal", "Profil, postuler, suivre, retirer", "Voir le pipeline interne d’une autre org"],
            ["Recruteur", "Principal", "Offres, pipeline, entretiens, décisions", "Voir les dossiers d’une autre organisation"],
            ["Owner org", "Principal", "Inviter membres, paramétrer le workspace", "— (tout voir dans son org)"],
            ["Hiring Manager", "Secondaire", "Scorecard / entretien sur dossiers assignés", "Publier une offre seul / gérer la facturation"],
            ["Administrateur", "Principal", "Vérif entreprises, supervision agrégée", "Lire le contenu nominatif croisé des orgs"],
            ["Scheduler", "Système", "Rappels SLA, digests", "Décider une embauche"],
            ["Matching local", "Système", "Calculer un score d’aide", "Rejeter ou accepter automatiquement"],
        ],
        [2.8, 2, 5.5, 5],
        size=8,
    )

    # ========== 5 FONCTIONNALITES CONVINCANTES ==========
    h(doc, "5. Identification des fonctionnalités")
    p(
        doc,
        "Chaque fonctionnalité est décrite ainsi : quoi · pour qui · pourquoi ça compte · "
        "comment on vérifie que c’est fait. Ce n’est pas une liste de boutons : c’est le contrat "
        "de livraison du produit.",
    )
    table(
        doc,
        ["ID", "Fonctionnalité (quoi)", "Pour qui", "Pourquoi c’est utile", "Preuve de done"],
        [
            [
                "F01",
                "Authentification email + OTP + JWT",
                "Tous",
                "Compte sûr, session contrôlée",
                "Login OK ; route protégée refuse sans token",
            ],
            [
                "F02",
                "Profil candidat + import CV (OCR local)",
                "Candidat",
                "Accélère la création de profil sans ressaisie",
                "Skills suggérées après upload",
            ],
            [
                "F03",
                "Catalogue d’offres publiques + recherche",
                "Candidat",
                "Trouver des postes pertinents rapidement",
                "Filtre / recherche renvoie les offres publiées",
            ],
            [
                "F04",
                "Candidature guidée (parcours multi-étapes)",
                "Candidat",
                "Dossier complet, moins d’abandons",
                "Candidature créée en statut SUBMITTED",
            ],
            [
                "F05",
                "Suivi bilatéral : statut + SLA restant",
                "Candidat",
                "Fini l’attente dans le noir",
                "Écran candidat montre étape + jours restants",
            ],
            [
                "F06",
                "CRUD / publication / clôture d’offre",
                "Recruteur",
                "Piloter le besoin ouvert",
                "Offre visible public seulement si publiée",
            ],
            [
                "F07",
                "SLA de réponse paramétrable sur l’offre",
                "Recruteur",
                "Engagement mesurable (anti-ghosting)",
                "Deadline = dépôt + N jours",
            ],
            [
                "F08",
                "Inbox + pipeline kanban priorisé",
                "Recruteur",
                "Traiter d’abord l’urgent (SLA / score)",
                "Tri visible ; retard mis en avant",
            ],
            [
                "F09",
                "Transition d’état côté serveur uniquement",
                "Recruteur",
                "Empêche la triche / saut d’étape",
                "Appel API illégal → 400/403",
            ],
            [
                "F10",
                "Rejet avec motif catalogue obligatoire",
                "Recruteur",
                "Justice du process + audit",
                "Rejet sans motif refusé",
            ],
            [
                "F11",
                "Planification d’entretiens typés",
                "Recruteur, HM",
                "Clarifier le rendez-vous (visio/présentiel…)",
                "Créneau + type enregistrés",
            ],
            [
                "F12",
                "Scorecard critères (tech / soft / culture)",
                "Recruteur, HM",
                "Évaluer de façon comparable",
                "Notes liées au dossier",
            ],
            [
                "F13",
                "Timeline immuable des transitions",
                "Les deux",
                "Preuve de ce qui s’est passé",
                "Historique append-only consultable",
            ],
            [
                "F14",
                "Matching local compétences (score)",
                "Les deux",
                "Prioriser la lecture humaine",
                "Score affiché ; aucun auto-rejet",
            ],
            [
                "F15",
                "IA locale bornée (suggestions, grille, language-check)",
                "Les deux",
                "Aide aux frictions réelles, sans cloud payant",
                "4 usages max ; décision humaine",
            ],
            [
                "F16",
                "Crédits IA Free / Premium",
                "Les deux",
                "Limiter l’abus, préparer la monétisation",
                "Compteur serveur ; refus si épuisé",
            ],
            [
                "F17",
                "Fiche entreprise + vérif ICE",
                "Recruteur, Admin",
                "Confiance sur l’émetteur de l’offre",
                "Badge vérifié après validation admin",
            ],
            [
                "F18",
                "Espace admin (orgs, comptes, indicateurs)",
                "Admin",
                "Gouverner la plateforme",
                "Actions admin isolées",
            ],
            [
                "F19",
                "Notifications + préférences",
                "Tous",
                "Informer sans noyer",
                "Événement → notif ; préférences respectées",
            ],
            [
                "F20",
                "Paramètres compte / apparence / plan",
                "Tous",
                "Confort et clarté du forfait",
                "Changements persistés",
            ],
            [
                "F21",
                "Export / droits données (esprit 09-08)",
                "Candidat",
                "Transparence et conformité",
                "Export accessible depuis le compte",
            ],
            [
                "F22",
                "Indicateur de fiabilité (org / candidat)",
                "Tous",
                "Rendre visible le respect des engagements",
                "Score calculé depuis la timeline",
            ],
        ],
        [1.1, 4.2, 2.2, 3.8, 3.5],
        size=7,
    )

    # ========== 6 REGLES METIER CLAIRES ==========
    h(doc, "6. Règles de métier")
    p(
        doc,
        "Une règle métier dit : quand telle situation arrive, le système doit faire ceci, "
        "parce que… Sinon, le produit devient incohérent ou trompeur. "
        "Toutes ces règles sont appliquées côté serveur (l’interface peut guider, elle ne décide pas).",
    )
    table(
        doc,
        ["ID", "Quand…", "Alors le système…", "Parce que…"],
        [
            [
                "R01",
                "On demande une transition d’état",
                "N’accepte que les sorties prévues par la machine à états",
                "Évite les sauts illégaux (ex. SUBMITTED → ACCEPTED)",
            ],
            [
                "R02",
                "Un recruteur agit sur une candidature",
                "Vérifie qu’elle appartient à son organisation / ses offres",
                "Empêche l’accès aux dossiers des autres (IDOR)",
            ],
            [
                "R03",
                "Un candidat consulte ou retire un dossier",
                "Autorise uniquement ses propres candidatures",
                "Confidentialité inter-candidats",
            ],
            [
                "R04",
                "Le statut cible est REJECTED",
                "Exige un motif du catalogue + enregistrement",
                "Pas de ghosting institutionnalisé ; audit possible",
            ],
            [
                "R05",
                "Une checklist d’étape est incomplète",
                "Bloque l’avance",
                "Le kanban ne remplace pas le process",
            ],
            [
                "R06",
                "Une candidature est déposée",
                "Calcule deadline = submitted_at + SLA de l’offre",
                "L’engagement de délai devient mesurable",
            ],
            [
                "R07",
                "Le matching / l’IA produit un score",
                "N’accepte / ne rejette jamais tout seul",
                "Décision finale = responsabilité humaine",
            ],
            [
                "R08",
                "Une action IA Free est demandée",
                "Consomme un crédit ; refuse si quota = 0",
                "Évite l’abus ; Premium géré côté serveur",
            ],
            [
                "R09",
                "Une route n’est pas publique",
                "Exige un JWT valide",
                "Pas d’API ouverte aux anonymes",
            ],
            [
                "R10",
                "Un mot de passe est stocké / un OTP émis",
                "Hash BCrypt ; OTP TTL + usage unique",
                "Secrets non réutilisables / non lisibles",
            ],
            [
                "R11",
                "Demande de retrait (WITHDRAWN)",
                "Autorise seulement le candidat propriétaire",
                "Un recruteur ne peut pas « effacer » un dossier à la place",
            ],
            [
                "R12",
                "Une transition réussit",
                "Ajoute une ligne d’historique (jamais d’effacement)",
                "Preuve temporelle du process",
            ],
            [
                "R13",
                "Une offre n’est pas publiée",
                "Elle n’apparaît pas au catalogue public",
                "Brouillons privés",
            ],
            [
                "R14",
                "Une entreprise est vérifiée",
                "Le statut vérifié est affiché sur ses offres",
                "Confiance asymétrique réduite côté candidat",
            ],
            [
                "R15",
                "Header Premium / rôle est spoofé côté client",
                "Le serveur recalcule les droits (ne fait pas confiance au front)",
                "Mentalité anti-triche",
            ],
        ],
        [1.1, 4.2, 5.2, 4.8],
        size=7,
    )

    h(doc, "6.1 Machine à états du pipeline", 2)
    table(
        doc,
        ["État actuel", "Signifie", "Peut aller vers"],
        [
            ["SUBMITTED", "Dossier reçu", "UNDER_REVIEW · WITHDRAWN"],
            ["UNDER_REVIEW", "En présélection", "INTERVIEW_SCHEDULED · REJECTED · WITHDRAWN"],
            ["INTERVIEW_SCHEDULED", "Entretien planifié", "INTERVIEW_COMPLETED · REJECTED · WITHDRAWN"],
            ["INTERVIEW_COMPLETED", "Entretien tenu", "ACCEPTED · REJECTED · INTERVIEW_SCHEDULED"],
            ["ACCEPTED", "Décision positive", "— (fin)"],
            ["REJECTED", "Décision négative motivée", "— (fin)"],
            ["WITHDRAWN", "Retrait candidat", "— (fin)"],
        ],
        [4, 5, 6.5],
        size=9,
    )

    # ========== 7 ==========
    h(doc, "7. Première réflexion sur le domaine")
    bullets(
        doc,
        [
            "L’objet central n’est ni « l’utilisateur » ni « l’offre » : c’est la candidature.",
            "Jobboards = acquisition ; HireLink = gouvernance. On ne réinvente pas LinkedIn.",
            "La bilatéralité est un choix produit : le candidat voit l’engagement (SLA, motif).",
            "L’IA est un accélérateur de lecture, pas un juge — sinon risque légal et éthique.",
            "Le front ment facilement : toute règle critique vit dans l’API.",
            "V1 = socle (auth, offres, pipeline sûr, matching, admin). Extensions ensuite.",
        ],
    )
    h(doc, "7.1 Hors scope V1", 2)
    bullets(
        doc,
        [
            "Paiement carte réel (Premium simulé en démo).",
            "LLM cloud payant.",
            "Remplacement d’un HRIS complet (paie, congés…).",
        ],
    )

    # ========== 8 CLASSES ==========
    h(doc, "8. Principales classes à utiliser")
    table(
        doc,
        ["Classe domaine", "Rôle", "Relations"],
        [
            ["User", "Compte authentifié", "1—* Application ; peut appartenir à org"],
            ["Organization", "Workspace entreprise", "1—* Membership ; 1—* JobOffer"],
            ["Membership", "Rôle dans une org", "*—1 User ; *—1 Organization"],
            ["JobOffer", "Offre + SLA", "*—1 Org ; 1—* Application"],
            ["Application", "Candidature / état", "*—1 Candidat ; *—1 Offre"],
            ["ApplicationStatusHistory", "Événement timeline", "*—1 Application"],
            ["Interview", "Entretien", "*—1 Application"],
            ["RefusalReason", "Motif catalogue", "référencé au rejet"],
            ["ApplicationRefusal", "Motif appliqué", "*—1 Application"],
            ["NotificationPreference", "Canaux utilisateur", "*—1 User"],
            ["AiCreditLedger", "Quota IA", "User ou Organization"],
            ["ReliabilityEvent", "Fait de fiabilité", "sujet User/Org"],
            ["MatchingAnalysis", "Score auditable", "lié Application"],
            ["ConsentRecord", "Consentement / export", "*—1 User"],
        ],
        [4, 4.5, 6.5],
        size=8,
    )
    h(doc, "8.1 Services applicatifs", 2)
    table(
        doc,
        ["Service", "Responsabilité"],
        [
            ["AuthService", "OTP, login, JWT"],
            ["JobOfferService", "Offres + ownership"],
            ["ApplicationService", "Postuler + transitions anti-triche"],
            ["SkillMatchingService", "Score Jaccard / couverture"],
            ["AiAssistService / AiCreditService", "IA bornée + crédits"],
            ["LanguageCheckService", "Formulation excluante"],
            ["ReliabilityService", "Scores de fiabilité"],
            ["NotificationService", "Événements → canaux"],
            ["CompanyVerificationService", "Vérif ICE / org"],
        ],
        [5.5, 10],
        size=9,
    )

    # ========== 9 BDD COMPLETE ==========
    doc.add_page_break()
    h(doc, "9. Conception de la base de données")
    p(
        doc,
        "SGBD cible : MySQL 8. Chaque table ci-dessous indique son utilité métier, "
        "puis ses colonnes avec le rôle de chaque colonne. Les types sont indicatifs "
        "(BIGINT, VARCHAR, TEXT, DATETIME, BOOLEAN, JSON, ENUM…).",
    )

    def table_block(title, utility, cols):
        h(doc, title, 2)
        p(doc, "Utilité : " + utility, size=10, after=3)
        table(doc, ["Colonne", "Type indicatif", "Utilité de la colonne"], cols, [4.2, 3.3, 8], size=8)

    table_block(
        "9.1 Table users",
        "Compte de connexion et profil de base (candidat, membre org, admin).",
        [
            ["id", "BIGINT PK", "Identifiant technique"],
            ["email", "VARCHAR UNIQUE", "Login ; contact"],
            ["password_hash", "VARCHAR", "Secret haché BCrypt (jamais en clair)"],
            ["full_name", "VARCHAR", "Nom affiché"],
            ["role", "ENUM", "CANDIDAT / RECRUTEUR / ADMINISTRATEUR"],
            ["skills", "TEXT", "Compétences candidat (matching)"],
            ["sought_roles", "VARCHAR", "Métiers ciblés"],
            ["preferred_locations", "VARCHAR", "Lieux préférés"],
            ["experience_level", "VARCHAR", "Junior / confirmé…"],
            ["company", "VARCHAR", "Nom entreprise (recruteur)"],
            ["company_ice", "VARCHAR", "Identifiant ICE"],
            ["company_verified", "BOOLEAN", "Statut vérification admin"],
            ["email_verified", "BOOLEAN", "Email confirmé (OTP)"],
            ["must_set_password", "BOOLEAN", "Force définition MDP"],
            ["created_at / updated_at", "DATETIME", "Audit"],
        ],
    )

    table_block(
        "9.2 Table organizations",
        "Workspace entreprise : racine multi-recruteurs.",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["name", "VARCHAR", "Raison sociale / nom commercial"],
            ["domain", "VARCHAR", "Domaine email pro (invitation)"],
            ["ice", "VARCHAR", "ICE officiel"],
            ["owner_user_id", "BIGINT FK", "Owner initial"],
            ["verified", "BOOLEAN", "Org vérifiée"],
            ["created_at", "DATETIME", "Création"],
        ],
    )

    table_block(
        "9.3 Table memberships",
        "Qui appartient à quelle organisation, avec quel rôle interne.",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["organization_id", "BIGINT FK", "Org"],
            ["user_id", "BIGINT FK", "Membre"],
            ["org_role", "ENUM", "OWNER / RECRUITER / HIRING_MANAGER"],
            ["invited_at", "DATETIME", "Invitation envoyée"],
            ["accepted_at", "DATETIME", "Acceptation"],
            ["active", "BOOLEAN", "Membre actif"],
        ],
    )

    table_block(
        "9.4 Table job_offers",
        "Offres d’emploi / stage / PFE avec engagement de délai.",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["organization_id", "BIGINT FK", "Org propriétaire"],
            ["recruiter_id", "BIGINT FK", "Auteur / responsable"],
            ["title", "VARCHAR", "Intitulé"],
            ["description", "TEXT", "Contenu de l’annonce"],
            ["required_skills", "TEXT", "Compétences pour matching"],
            ["contract_type", "VARCHAR", "CDI / stage / PFE…"],
            ["location", "VARCHAR", "Lieu"],
            ["response_sla_days", "INT", "Engagement de réponse (jours)"],
            ["published", "BOOLEAN", "Visible catalogue public"],
            ["closed", "BOOLEAN", "Offre clôturée"],
            ["created_at / updated_at", "DATETIME", "Audit"],
        ],
    )

    table_block(
        "9.5 Table applications",
        "Candidature = objet central du pipeline.",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["candidate_id", "BIGINT FK", "Candidat"],
            ["job_offer_id", "BIGINT FK", "Offre"],
            ["status", "ENUM", "État courant machine à états"],
            ["match_score", "DECIMAL NULL", "Aide au tri (nullable)"],
            ["cover_letter", "TEXT", "Lettre optionnelle"],
            ["submitted_at", "DATETIME", "Point de départ du SLA"],
            ["updated_at", "DATETIME", "Dernière transition"],
            ["interview_type", "VARCHAR NULL", "Type entretien courant"],
            ["interview_at", "DATETIME NULL", "Créneau"],
            ["UNIQUE(candidate_id, job_offer_id)", "contrainte", "Une candidature par offre"],
        ],
    )

    table_block(
        "9.6 Table application_status_history",
        "Timeline immuable : chaque transition = une ligne (append-only).",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["application_id", "BIGINT FK", "Dossier concerné"],
            ["from_status", "VARCHAR", "État avant"],
            ["to_status", "VARCHAR", "État après"],
            ["actor_user_id", "BIGINT FK", "Qui a agi"],
            ["payload_json", "JSON", "Motif, checklist, notes techniques"],
            ["created_at", "DATETIME", "Horodatage (preuve)"],
        ],
    )

    table_block(
        "9.7 Table interviews",
        "Entretiens liés à une candidature.",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["application_id", "BIGINT FK", "Dossier"],
            ["type", "ENUM", "PHONE / VISIO / ONSITE / ASYNC"],
            ["scheduled_at", "DATETIME", "Date/heure"],
            ["location_or_link", "VARCHAR", "Lieu ou URL"],
            ["status", "VARCHAR", "PLANNED / DONE / NO_SHOW / CANCELED"],
            ["created_by", "BIGINT FK", "Organisateur"],
        ],
    )

    table_block(
        "9.8 Tables refusal_reason_catalog & application_refusals",
        "Catalogue de motifs + motif appliqué à un rejet (R04).",
        [
            ["refusal_reason_catalog.id", "BIGINT PK", "Motif"],
            ["code", "VARCHAR UNIQUE", "Code stable (ex. PROFILE_MISMATCH)"],
            ["label", "VARCHAR", "Libellé affiché"],
            ["active", "BOOLEAN", "Toujours proposé ?"],
            ["application_refusals.id", "BIGINT PK", "Instance de rejet"],
            ["application_id", "BIGINT FK", "Dossier rejeté"],
            ["reason_code", "VARCHAR FK", "Motif choisi"],
            ["comment", "VARCHAR", "Commentaire borné optionnel"],
            ["created_at", "DATETIME", "Quand"],
        ],
    )

    table_block(
        "9.9 Tables notifications",
        "Préférences utilisateur + journal des envois.",
        [
            ["notification_preferences.user_id", "BIGINT FK", "Utilisateur"],
            ["channel", "ENUM", "IN_APP / EMAIL / PUSH / WHATSAPP"],
            ["enabled", "BOOLEAN", "Canal activé"],
            ["notification_deliveries.id", "BIGINT PK", "Envoi"],
            ["event_type", "VARCHAR", "STATUS_CHANGED / SLA_WARNING…"],
            ["payload_json", "JSON", "Contenu"],
            ["sent_at", "DATETIME", "Horodatage"],
            ["digest_bucket", "VARCHAR NULL", "Regroupement digest"],
        ],
    )

    table_block(
        "9.10 Table ai_credit_ledgers",
        "Compteur mensuel des crédits IA (anti-abus).",
        [
            ["id", "BIGINT PK", "Identifiant"],
            ["subject_type", "ENUM", "USER / ORGANIZATION"],
            ["subject_id", "BIGINT", "Qui consomme"],
            ["period_ym", "CHAR(7)", "Mois (2026-09)"],
            ["used", "INT", "Crédits consommés"],
            ["limit_free", "INT", "Plafond Free"],
            ["action", "VARCHAR", "PROFILE_REWRITE / INTERVIEW_KIT…"],
            ["created_at", "DATETIME", "Mouvement"],
        ],
    )

    table_block(
        "9.11 Tables reliability_events & reliability_badges",
        "Faits pour le score de fiabilité + badge public optionnel.",
        [
            ["reliability_events.id", "BIGINT PK", "Événement"],
            ["subject_type", "ENUM", "ORG / CANDIDATE"],
            ["subject_id", "BIGINT", "Sujet"],
            ["event_type", "VARCHAR", "SLA_OK / EXPLICIT_DECISION / NO_SHOW…"],
            ["weight", "DECIMAL", "Poids dans le calcul"],
            ["created_at", "DATETIME", "Quand"],
            ["reliability_badges.organization_id", "BIGINT FK", "Org"],
            ["token", "VARCHAR", "Intégration page carrière"],
            ["score_snapshot", "INT", "Score affiché"],
            ["view_count", "INT", "Affichages (PLG)"],
        ],
    )

    table_block(
        "9.12 Tables complémentaires",
        "Matching auditable, consentements, parrainage, language-check, panélistes.",
        [
            ["matching_analyses", "table", "application_id, score, details_json, created_at"],
            ["consent_records", "table", "user_id, purpose, granted, created_at"],
            ["referral_links", "table", "code, referrer_user_id, status, bonus_granted"],
            ["job_offer_language_checks", "table", "offer_id, markers_json, accepted, checked_at"],
            ["panelist_access_links", "table", "token, application_id, expires_at, scopes"],
            ["email_otps", "table", "email, code_hash, expires_at, used"],
        ],
    )

    p(
        doc,
        "Relations principales : Organization —< Membership >— User ; "
        "Organization —< JobOffer —< Application >— User(Candidat) ; "
        "Application —< History / Interviews / Refusal.",
        size=10,
        italic=True,
    )

    # ========== ANNEXE A ==========
    doc.add_page_break()
    h(doc, "Annexe A — Benchmarking marché")
    p(
        doc,
        "Sources de cadrage : AfricaWork (Top platforms Morocco 2025 ; LinkedIn / jobboards / ATS), "
        "guides ATS Maroc 2026, comparatifs internationaux. Lecture : le volume est couvert ; "
        "HireLink se différencie sur la gouvernance bilatérale (SLA + motif + timeline + matching local).",
        size=10,
    )
    table(
        doc,
        ["Plateforme", "Rôle marché", "Force", "Ce qui manque vs HireLink"],
        [
            ["ReKrute", "Jobboard MA", "Visibilité nationale", "Pas de pipeline / SLA / motif"],
            ["Emploi.ma", "Jobboard MA", "Volume multi-secteurs", "Pas d’orchestration ATS"],
            ["Dreamjob", "Jobboard junior", "Jeunes diplômés", "Gouvernance faible"],
            ["Indeed Maroc", "Agrégateur", "Diffusion large", "Pas bilatéral"],
            ["LinkedIn Recruiter", "Sourcing", "Profils qualifiés", "Coût ; pipeline limité hors offre entreprise"],
            ["Bayt", "MENA", "Couverture régionale", "Peu adapté PME stage–PFE MA"],
            ["Greenhouse / Lever", "ATS US", "Scorecard / CRM", "Prix / complexité ; asymétrie confiance"],
            ["Teamtailor / Workable", "ATS SaaS", "UX / équilibre", "Pas d’engagement de délai public"],
            ["Workday / Taleo", "Grands comptes", "Industrialisation", "Hors PME légère"],
            ["HRMAPS / locaux", "ATS MA", "Ancrage local", "SLA visibles + motifs + matching local borné"],
            ["HireLink", "ATS bilatéral", "SLA · motif · timeline · matching local · sécurité API", "—"],
        ],
        [3.2, 2.8, 4, 5.5],
        size=8,
    )

    h(doc, "Annexe B — Sources", 1)
    bullets(
        doc,
        [
            "AfricaWork — Ranking Top 7 Recruitment Platforms in Morocco 2025",
            "AfricaWork — Recruitment Technology in Morocco: LinkedIn, Job Boards, and ATS Compared (LinkedIn > 7 M users MA)",
            "SHRM — employer ghosting (~60 % candidats ; 72 % en parlent)",
            "LinkedIn Global Talent Trends — 75 % jugent l’entreprise via le process de recrutement",
            "CNDP / Loi 09-08 — protection des données à caractère personnel (Maroc)",
            "Bonnes pratiques SLA recrutement (communication candidat, anti-silence)",
        ],
    )

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.paragraph_format.space_before = Pt(18)
    rs(end.add_run("— Fin du cahier des charges HireLink —"), italic=True, size=10, color=ROYAL)

    paths = []
    for d in OUT:
        d.mkdir(parents=True, exist_ok=True)
        for fname in (NAME, "HireLink_Cahier_des_charges_v4.docx"):
            path = d / fname
            try:
                doc.save(path)
                paths.append(path)
            except PermissionError:
                alt = d / fname.replace(".docx", "_CLAIR.docx")
                doc.save(alt)
                paths.append(alt)
                print("LOCKED ->", alt)
    return paths


if __name__ == "__main__":
    for path in build():
        print(path, path.stat().st_size)
